import threading
import time
import os
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from sift_indexer.api import create_app
from sift_indexer.db import Database, SCHEMA_VERSION
from sift_indexer.indexer import Indexer
from sift_indexer import indexer as indexer_module
from sift_indexer import parsers as parser_module
from sift_indexer.parsers import IsolatedParser, ParseTimeoutError, parse


def run_scan(indexer: Indexer, root_id: int):
    job_id = indexer.start_scan(root_id)
    assert job_id is not None
    deadline = time.monotonic() + 45
    while time.monotonic() < deadline:
        job = indexer.job(job_id)
        if job["state"] in {"done", "cancelled", "failed"}:
            assert job["state"] == "done", job
            return job
        time.sleep(0.05)
    raise AssertionError(f"Scan deadline exceeded: {indexer.job(job_id)!r}")


def test_incremental_identity_missing_duplicates_and_errors(tmp_path: Path):
    library = tmp_path / "library"
    library.mkdir()
    (library / "notes.txt").write_text("automata state machine", encoding="utf-8")
    (library / "broken.pdf").write_bytes(b"not a PDF")
    (library / "other.bin").write_bytes(b"ignored")
    docx = Document()
    docx.add_paragraph("document text")
    docx.save(library / "word.docx")
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with (library / "blank.pdf").open("wb") as stream:
        writer.write(stream)

    db = Database(tmp_path / "data" / "index.sqlite")
    indexer = Indexer(db)
    root_id = indexer.add_root(str(library))["id"]
    first = run_scan(indexer, root_id)
    assert (first["created"], first["failed"], first["skipped"]) == (3, 1, 1)
    with db.connect() as connection:
        original_id = connection.execute("SELECT id FROM files WHERE filename='notes.txt'").fetchone()[0]
        assert connection.execute("SELECT COUNT(*) FROM documents_fts WHERE documents_fts MATCH 'automata'").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM parser_warnings").fetchone()[0] == 1

    second = run_scan(indexer, root_id)
    assert second["created"] == 0 and second["unchanged"] == 3 and second["failed"] == 1
    (library / "notes.txt").rename(library / "renamed.txt")
    moved = run_scan(indexer, root_id)
    assert moved["moved"] == 1
    with db.connect() as connection:
        assert connection.execute("SELECT id FROM files WHERE filename='renamed.txt'").fetchone()[0] == original_id
        assert connection.execute("SELECT COUNT(*) FROM path_history WHERE file_id=?", (original_id,)).fetchone()[0] == 1

    (library / "copy.txt").write_text("automata state machine", encoding="utf-8")
    duplicate = run_scan(indexer, root_id)
    assert duplicate["created"] == 1
    with db.connect() as connection:
        assert connection.execute("SELECT duplicate_of FROM files WHERE filename='copy.txt'").fetchone()[0] == original_id

    (library / "renamed.txt").write_text("changed content", encoding="utf-8")
    changed = run_scan(indexer, root_id)
    assert changed["changed"] == 1
    with db.connect() as connection:
        assert connection.execute("SELECT duplicate_of FROM files WHERE filename='copy.txt'").fetchone()[0] is None
        assert connection.execute("SELECT COUNT(*) FROM documents_fts WHERE documents_fts MATCH 'automata'").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM documents_fts WHERE documents_fts MATCH 'changed'").fetchone()[0] == 1
    (library / "renamed.txt").unlink()
    missing = run_scan(indexer, root_id)
    assert missing["missing"] == 1
    with db.connect() as connection:
        assert connection.execute("SELECT status FROM files WHERE id=?", (original_id,)).fetchone()[0] == "missing"


def test_root_removal_retains_cache_and_symlink_is_skipped(tmp_path: Path):
    library = tmp_path / "library"
    library.mkdir()
    outside = tmp_path / "outside.txt"
    outside.write_text("secret", encoding="utf-8")
    (library / "inside.txt").write_text("safe", encoding="utf-8")
    try:
        (library / "link.txt").symlink_to(outside)
    except OSError:
        pass  # Windows systems without developer mode cannot create test symlinks.
    db = Database(tmp_path / "index.sqlite")
    indexer = Indexer(db)
    root_id = indexer.add_root(str(library))["id"]
    job = run_scan(indexer, root_id)
    assert job["created"] == 1
    assert job["skipped"] == (1 if (library / "link.txt").is_symlink() else 0)
    indexer.set_root_state(root_id, "removed")
    with db.connect() as connection:
        assert connection.execute("SELECT status FROM files").fetchone()[0] == "excluded"
        assert connection.execute("SELECT text FROM documents").fetchone()[0] == "safe"


def test_private_api_search_and_restart(tmp_path: Path):
    library = tmp_path / "library"
    library.mkdir()
    (library / "lesson.txt").write_text("finite state automata", encoding="utf-8")
    named_file = Document()
    named_file.core_properties.title = "Different title"
    named_file.add_paragraph("unrelated text")
    named_file.save(library / "course-guide.docx")
    db_path = tmp_path / "data" / "library.sqlite"
    token = {"X-SIFT-Token": "test-secret"}
    with TestClient(create_app(db_path, "test-secret")) as client:
        assert client.get("/health").status_code == 401
        assert client.get("/health", headers=token).json()["schema_version"] == SCHEMA_VERSION
        root = client.post("/folders", headers=token, json={"path": str(library)}).json()
        job_id = client.post("/scans", headers=token, json={"root_id": root["id"]}).json()["id"]
        for _ in range(100):
            job = client.get(f"/jobs/{job_id}", headers=token).json()
            if job["state"] == "done":
                break
            time.sleep(0.05)
        assert job["state"] == "done"
        results = client.post("/search", headers=token, json={"query": "automata"}).json()
        assert results['total'] == 1 and results['items'][0]["filename"] == "lesson.txt"
        filename_results = client.post("/search", headers=token, json={"query": "course-guide"}).json()
        assert any(item["filename"] == "course-guide.docx" for item in filename_results['items'])
    with TestClient(create_app(db_path, "test-secret")) as client:
        assert client.get("/documents", headers=token).json()["total"] == 2


def test_cancel_does_not_publish_partial_document(tmp_path: Path, monkeypatch):
    library = tmp_path / "library"
    library.mkdir()
    (library / "slow.txt").write_text("study text", encoding="utf-8")
    indexer = Indexer(Database(tmp_path / "index.sqlite"))
    root_id = indexer.add_root(str(library))["id"]
    started = threading.Event()
    release = threading.Event()
    original_parse = indexer_module.parse

    def slow_parse(path):
        started.set()
        assert release.wait(5)
        return original_parse(path)

    monkeypatch.setattr(indexer_module, "parse", slow_parse)
    job_id = indexer.start_scan(root_id)
    assert started.wait(5)
    assert indexer.cancel(job_id)
    release.set()
    for _ in range(100):
        if indexer.job(job_id)["state"] == "cancelled":
            break
        time.sleep(0.05)
    assert indexer.job(job_id)["state"] == "cancelled"
    with indexer.db.connect() as connection:
        assert connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0


def test_watcher_indexes_new_file(tmp_path: Path):
    library = tmp_path / "library"
    library.mkdir()
    token = {"X-SIFT-Token": "test-secret"}
    with TestClient(create_app(tmp_path / "data" / "index.sqlite", "test-secret")) as client:
        client.post("/folders", headers=token, json={"path": str(library)})
        (library / "new.txt").write_text("watched content", encoding="utf-8")
        for _ in range(120):
            if client.get("/documents", headers=token).json()["total"] == 1:
                break
            time.sleep(0.05)
        assert client.get("/documents", headers=token).json()["total"] == 1
        assert client.get("/jobs/latest", headers=token).json()["source"] == "watcher"


def test_utf16_txt_and_binary_txt(tmp_path: Path):
    text_file = tmp_path / "unicode.txt"
    text_file.write_text("study notes", encoding="utf-16")
    assert parse(text_file).metadata["encoding"] == "utf-16"
    binary_file = tmp_path / "binary.txt"
    binary_file.write_bytes(b"\x00\x01\x02")
    try:
        parse(binary_file)
        raise AssertionError("Binary TXT should be rejected")
    except ValueError as exc:
        assert "binary" in str(exc)


def test_case_only_rename_keeps_identity_on_windows(tmp_path: Path):
    if os.name != "nt":
        return
    library = tmp_path / "library"
    library.mkdir()
    original = library / "Case.txt"
    original.write_text("same content", encoding="utf-8")
    db = Database(tmp_path / "index.sqlite")
    indexer = Indexer(db)
    root_id = indexer.add_root(str(library))["id"]
    run_scan(indexer, root_id)
    original.rename(library / "case.txt")
    job = run_scan(indexer, root_id)
    assert job["moved"] == 1
    with db.connect() as connection:
        assert connection.execute("SELECT COUNT(*) FROM files").fetchone()[0] == 1
        assert connection.execute("SELECT filename FROM files").fetchone()[0] == "case.txt"


def test_scan_requested_during_scan_is_queued_and_reconciled(tmp_path: Path, monkeypatch):
    library = tmp_path / "library"
    library.mkdir()
    (library / "first.txt").write_text("first", encoding="utf-8")
    db = Database(tmp_path / "index.sqlite")
    indexer = Indexer(db)
    root_id = indexer.add_root(str(library))["id"]
    entered = threading.Event()
    release = threading.Event()
    original_parse = indexer_module.parse

    def wait_once(path):
        if path.name == "first.txt" and not entered.is_set():
            entered.set()
            assert release.wait(5)
        return original_parse(path)

    monkeypatch.setattr(indexer_module, "parse", wait_once)
    first_id = indexer.start_scan(root_id)
    assert entered.wait(5)
    (library / "second.txt").write_text("second", encoding="utf-8")
    second_id = indexer.start_scan(root_id)
    assert second_id != first_id
    assert indexer.job(second_id)["state"] == "queued"
    release.set()
    for _ in range(120):
        with db.connect() as connection:
            count = connection.execute("SELECT COUNT(*) FROM files WHERE status='indexed'").fetchone()[0]
        if count == 2 and indexer.job(second_id)["state"] == "done":
            break
        time.sleep(0.05)
    assert count == 2
    assert indexer.job(first_id)["state"] == "done"
    assert indexer.job(second_id)["state"] == "done"


def test_cancelling_queued_scan_does_not_cancel_running_scan(tmp_path: Path, monkeypatch):
    library = tmp_path / "library"
    library.mkdir()
    (library / "first.txt").write_text("first", encoding="utf-8")
    indexer = Indexer(Database(tmp_path / "index.sqlite"))
    root_id = indexer.add_root(str(library))["id"]
    entered = threading.Event()
    release = threading.Event()
    original_parse = indexer_module.parse

    def wait_parse(path):
        if path.name == "first.txt" and not entered.is_set():
            entered.set()
            assert release.wait(5)
        return original_parse(path)

    monkeypatch.setattr(indexer_module, "parse", wait_parse)
    running_id = indexer.start_scan(root_id)
    assert entered.wait(5)
    queued_id = indexer.start_scan(root_id)
    assert indexer.job(queued_id)["state"] == "queued"
    assert indexer.cancel(queued_id)
    assert indexer.job(queued_id)["state"] == "cancelled"
    assert indexer.job(running_id)["state"] == "running"
    release.set()
    for _ in range(100):
        if indexer.job(running_id)["state"] == "done":
            break
        time.sleep(0.05)
    assert indexer.job(running_id)["state"] == "done"
    assert indexer.job(queued_id)["state"] == "cancelled"


def test_pausing_one_root_does_not_stop_other_roots(tmp_path: Path, monkeypatch):
    first = tmp_path / "first"
    second = tmp_path / "second"
    first.mkdir(); second.mkdir()
    (first / "slow.txt").write_text("first root", encoding="utf-8")
    (second / "other.txt").write_text("second root", encoding="utf-8")
    db = Database(tmp_path / "index.sqlite")
    indexer = Indexer(db)
    first_id = indexer.add_root(str(first))["id"]
    second_id = indexer.add_root(str(second))["id"]
    entered = threading.Event()
    release = threading.Event()
    original_parse = indexer_module.parse

    def pause_first(path):
        if path.name == "slow.txt":
            entered.set()
            assert release.wait(5)
        return original_parse(path)

    monkeypatch.setattr(indexer_module, "parse", pause_first)
    job_id = indexer.start_scan()
    assert entered.wait(5)
    indexer.set_root_state(first_id, "paused")
    release.set()
    for _ in range(120):
        if indexer.job(job_id)["state"] in {"done", "cancelled", "failed"}:
            break
        time.sleep(0.05)
    assert indexer.job(job_id)["state"] == "done"
    with db.connect() as connection:
        assert connection.execute("SELECT status FROM files WHERE root_id=?", (first_id,)).fetchone() is None
        assert connection.execute("SELECT status FROM files WHERE root_id=?", (second_id,)).fetchone()[0] == "indexed"


def test_parser_rejects_oversized_input_before_reading(tmp_path: Path, monkeypatch):
    source = tmp_path / "large.txt"
    source.write_bytes(b"123456")
    monkeypatch.setattr(parser_module, "MAX_INPUT_BYTES", 5)
    try:
        parse(source)
        raise AssertionError("Oversized input should be rejected")
    except parser_module.ParseLimitError as exc:
        assert "extraction limit" in str(exc)


def test_parser_timeout_and_cancellation_stop_worker(tmp_path: Path):
    source = tmp_path / "note.txt"
    source.write_text("safe to parse", encoding="utf-8")
    parser = IsolatedParser()
    try:
        try:
            parser.parse(source, lambda: False, timeout=0.000001)
            raise AssertionError("Expired parsing budget should stop the worker")
        except ParseTimeoutError:
            assert parser._process is None

        checks = 0
        def stop_during_parse():
            nonlocal checks
            checks += 1
            return checks >= 2

        try:
            parser.parse(source, stop_during_parse)
            raise AssertionError("Cancellation should stop the worker")
        except InterruptedError:
            assert parser._process is None
    finally:
        parser.close()


def test_folder_can_be_paused_and_resumed_through_private_api(tmp_path: Path):
    library = tmp_path / "library"
    library.mkdir()
    token = {"X-SIFT-Token": "test-secret"}
    with TestClient(create_app(tmp_path / "data" / "index.sqlite", "test-secret")) as client:
        root = client.post("/folders", headers=token, json={"path": str(library)}).json()
        paused = client.patch(f"/folders/{root['id']}", headers=token, json={"state": "paused"})
        assert paused.status_code == 200
        resumed = client.patch(f"/folders/{root['id']}", headers=token, json={"state": "active"})
        assert resumed.status_code == 200
