"""Local parsers with a shared result shape. Source files are opened read-only."""

from __future__ import annotations

import multiprocessing
import threading
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader


PARSER_VERSION = 1
SUPPORTED = {".pdf", ".docx", ".txt"}
MAX_INPUT_BYTES = 64 * 1024 * 1024
MAX_DOCX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_DOCX_ENTRIES = 20_000
MAX_PDF_PAGES = 2_000
MAX_EXTRACTED_CHARS = 16_000_000
PARSER_TIMEOUT_SECONDS = 30.0


class ParseLimitError(ValueError):
    """The source exceeds an explicit, user-visible extraction limit."""


class ParseTimeoutError(TimeoutError):
    """Parsing exceeded its wall-clock budget."""


class IsolatedParser:
    """One reusable parser process, replaceable immediately on cancel or timeout."""

    def __init__(self):
        self._guard = threading.Lock()
        self._process: multiprocessing.Process | None = None
        self._connection = None

    def _start(self) -> None:
        if self._process is not None and self._process.is_alive():
            return
        if self._process is not None:
            self._stop_locked()
        context = multiprocessing.get_context("spawn")
        parent, child = context.Pipe(duplex=True)
        process = context.Process(target=_parser_worker, args=(child,), daemon=True)
        process.start()
        child.close()
        self._process = process
        self._connection = parent

    def parse(self, path: Path, should_stop, timeout: float = PARSER_TIMEOUT_SECONDS) -> "ParsedDocument":
        with self._guard:
            if should_stop():
                raise InterruptedError("Scan stopped before parsing")
            deadline = time.monotonic() + timeout
            self._start()
            process, connection = self._process, self._connection
            assert process is not None and connection is not None
            try:
                if time.monotonic() >= deadline:
                    raise ParseTimeoutError(f"Exceeded the {timeout:g} second parsing limit")
                connection.send((str(path), timeout))
                while True:
                    if should_stop():
                        raise InterruptedError("Scan stopped during parsing")
                    remaining = deadline - time.monotonic()
                    if remaining <= 0:
                        raise ParseTimeoutError(f"Exceeded the {timeout:g} second parsing limit")
                    if connection.poll(min(0.05, remaining)):
                        if time.monotonic() > deadline:
                            raise ParseTimeoutError(f"Exceeded the {timeout:g} second parsing limit")
                        response = connection.recv()
                        if response[0] == "ok":
                            return response[1]
                        _, module_name, type_name, message = response
                        detail = f"{type_name}: {message}"
                        if module_name == __name__ and type_name == "ParseTimeoutError":
                            raise ParseTimeoutError(message)
                        if module_name == "builtins" and type_name == "TimeoutError":
                            raise ParseTimeoutError(detail)
                        if module_name == __name__ and type_name == "ParseLimitError":
                            raise ParseLimitError(message)
                        raise ValueError(detail)
                    if not process.is_alive():
                        raise RuntimeError(f"Parser worker stopped unexpectedly (exit {process.exitcode})")
            except (InterruptedError, ParseTimeoutError, OSError, EOFError, RuntimeError):
                self._stop_locked()
                raise

    def _stop_locked(self) -> None:
        process, connection = self._process, self._connection
        self._process = None
        self._connection = None
        if connection is not None:
            try:
                connection.close()
            except OSError:
                pass
        if process is not None:
            if process.is_alive():
                process.terminate()
            process.join(timeout=2)
            if process.is_alive() and hasattr(process, "kill"):
                process.kill()
                process.join(timeout=2)
            try:
                process.close()
            except (OSError, ValueError):
                pass

    def close(self) -> None:
        with self._guard:
            process, connection = self._process, self._connection
            if process is None:
                return
            try:
                if process.is_alive() and connection is not None:
                    connection.send(None)
                    process.join(timeout=0.5)
            except (OSError, EOFError, BrokenPipeError):
                pass
            self._stop_locked()


def _parser_worker(connection) -> None:
    try:
        while True:
            request = connection.recv()
            if request is None:
                return
            raw_path, timeout = request
            try:
                document = parse_document(Path(raw_path), time.monotonic() + timeout)
                connection.send(("ok", document))
            except Exception as exc:
                connection.send(("error", type(exc).__module__, type(exc).__name__, str(exc)[:1500]))
    except (EOFError, OSError, BrokenPipeError):
        return
    finally:
        connection.close()


@dataclass
class TextBlock:
    position: int
    location: str | None
    text: str


@dataclass
class ParsedDocument:
    title: str
    text: str
    blocks: list[TextBlock] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


def _check_deadline(deadline: float | None) -> None:
    if deadline is not None and time.monotonic() >= deadline:
        raise ParseTimeoutError(f"Exceeded the {PARSER_TIMEOUT_SECONDS:g} second parsing limit")


def _append_block(blocks: list[TextBlock], location: str | None, text: str,
                  deadline: float | None, current_chars: int) -> int:
    _check_deadline(deadline)
    total = current_chars + len(text)
    if total > MAX_EXTRACTED_CHARS:
        raise ParseLimitError(f"Extracted text exceeds the {MAX_EXTRACTED_CHARS:,} character limit")
    blocks.append(TextBlock(len(blocks), location, text))
    return total


def parse_document(path: Path, deadline: float | None = None) -> ParsedDocument:
    try:
        size = path.stat().st_size
    except OSError:
        raise
    if size > MAX_INPUT_BYTES:
        raise ParseLimitError(f"File exceeds the {MAX_INPUT_BYTES // (1024 * 1024)} MB extraction limit")
    suffix = path.suffix.lower()
    if suffix == ".txt":
        _check_deadline(deadline)
        with path.open("rb") as stream:
            data = stream.read(MAX_INPUT_BYTES + 1)
        if len(data) > MAX_INPUT_BYTES:
            raise ParseLimitError(f"File exceeds the {MAX_INPUT_BYTES // (1024 * 1024)} MB extraction limit")
        _check_deadline(deadline)
        if data.startswith((b"\xff\xfe", b"\xfe\xff")):
            content = data.decode("utf-16")
            encoding = "utf-16"
        else:
            if b"\x00" in data:
                raise ValueError("TXT file contains binary data")
            try:
                content = data.decode("utf-8-sig")
                encoding = "utf-8"
            except UnicodeDecodeError:
                content = data.decode("cp1252")
                encoding = "cp1252"
        if len(content) > MAX_EXTRACTED_CHARS:
            raise ParseLimitError(f"Extracted text exceeds the {MAX_EXTRACTED_CHARS:,} character limit")
        _check_deadline(deadline)
        warnings = [] if content.strip() else ["No extractable text"]
        return ParsedDocument(path.stem, content, [TextBlock(0, None, content)] if content else [], warnings, {"encoding": encoding})
    if suffix == ".pdf":
        _check_deadline(deadline)
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            if not reader.decrypt(""):
                raise ValueError("Password-protected PDF")
        blocks: list[TextBlock] = []
        warnings: list[str] = []
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ParseLimitError(f"PDF exceeds the {MAX_PDF_PAGES:,} page extraction limit")
        text_length = 0
        for number, page in enumerate(reader.pages, 1):
            _check_deadline(deadline)
            try:
                content = page.extract_text() or ""
            except Exception as exc:
                warnings.append(f"Page {number}: {type(exc).__name__}: {exc}")
                continue
            if content.strip():
                text_length += len(content)
                if text_length > MAX_EXTRACTED_CHARS:
                    raise ParseLimitError(f"Extracted text exceeds the {MAX_EXTRACTED_CHARS:,} character limit")
                blocks.append(TextBlock(number - 1, f"Page {number}", content))
            else:
                warnings.append(f"Page {number}: no extractable text (possibly scanned)")
        title = str(reader.metadata.title).strip() if reader.metadata and reader.metadata.title else path.stem
        return ParsedDocument(title, "\n\n".join(block.text for block in blocks), blocks, warnings,
                              {"pages": len(reader.pages)})
    if suffix == ".docx":
        _check_deadline(deadline)
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_DOCX_ENTRIES:
                    raise ParseLimitError(f"DOCX exceeds the {MAX_DOCX_ENTRIES:,} archive-entry limit")
                expanded_bytes = sum(entry.file_size for entry in entries)
                if expanded_bytes > MAX_DOCX_EXPANDED_BYTES:
                    raise ParseLimitError("DOCX expands beyond the 256 MB safety limit")
        except zipfile.BadZipFile:
            raise ValueError("Invalid DOCX archive") from None
        document = WordDocument(str(path))
        blocks = []
        text_length = 0
        for paragraph in document.paragraphs:
            _check_deadline(deadline)
            content = paragraph.text.strip()
            if content:
                text_length = _append_block(blocks, f"Paragraph {len(blocks) + 1}", content, deadline, text_length)
        for table_number, table in enumerate(document.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                _check_deadline(deadline)
                content = " | ".join(cell.text.strip() for cell in row.cells)
                if content.strip(" |"):
                    text_length = _append_block(blocks, f"Table {table_number}, row {row_number}", content, deadline, text_length)
        title = document.core_properties.title.strip() if document.core_properties.title else path.stem
        title = title or path.stem
        warnings = [] if blocks else ["No extractable text"]
        return ParsedDocument(title, "\n\n".join(block.text for block in blocks), blocks, warnings)
    raise ValueError(f"Unsupported file type: {suffix}")


def parse(path: Path) -> ParsedDocument:
    """Convenience parser for CLI/tests; scans use the isolated, timed worker."""
    return parse_document(path)
